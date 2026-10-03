import os
import re
import time
import io
import logging
import shutil
import tempfile
from importlib.util import find_spec
from typing import Optional, Tuple, Dict, Any, List
from pathlib import Path

import requests
from fastapi import UploadFile
from PIL import Image

from .connection_registry import registry
from .audit import log_action

# MORE Lazy imports for heavy libs
# import numpy as np
# import cv2
# from pdf2image import convert_from_path

_paddle_instance = None
_OCR_ENGINE = os.getenv("OCR_ENGINE", "tesseract").lower()
_OCR_MIN_TEXT_CHARS = int(os.getenv("OCR_MIN_TEXT_CHARS", "200"))
_OCR_ENGINE_TTL_SEC = int(os.getenv("OCR_ENGINE_TTL_SEC", "900"))

_paddle_engine: Optional[object] = None
_paddle_last_used: float = 0.0

logger = logging.getLogger(__name__)

# Health check OCRs this bundled image and expects the token below in the output.
_HEALTH_IMAGE = Path(__file__).resolve().parents[1] / "assets" / "ocr_health_check.png"
_HEALTH_EXPECTED_TOKEN = "2468"
_HEALTH_TTL_SEC = int(os.getenv("OCR_HEALTH_TTL_SEC", "600"))
_health_cache: Optional[Tuple[float, Dict[str, Any]]] = None


def _now() -> float:
    return time.time()


def _normalize_engine_name(value: Optional[str]) -> str:
    """Normalize configured OCR engine aliases without changing valid values."""
    normalized = (value or "").strip().lower().replace("_", "").replace("-", "")
    if normalized in {"paddle", "paddleocr"}:
        return "paddle"
    if normalized in {"tesseract", "pytesseract"}:
        return "tesseract"
    return normalized or "tesseract"


def _resolve_engine() -> str:
    connector = registry.get("ocr-default") or next(
        (c for c in registry.list("ocr").values()), None
    )
    if connector:
        engine = connector.metadata.get("engine")
        if engine:
            return _normalize_engine_name(str(engine))
    return _normalize_engine_name(_OCR_ENGINE)


def _should_unload(last_used: float) -> bool:
    if not last_used:
        return False
    return (_now() - last_used) > _OCR_ENGINE_TTL_SEC


def get_paddle() -> Tuple[Optional[object], Optional[str]]:
    global _paddle_engine, _paddle_last_used
    if _paddle_engine is not None and _should_unload(_paddle_last_used):
        _paddle_engine = None
    if _paddle_engine is not None:
        return _paddle_engine, None
    try:
        from paddleocr import PaddleOCR  # type: ignore
        # use_angle_cls=True loads lighter model?
        # lang='en'
        _paddle_engine = PaddleOCR(use_angle_cls=True, lang="en", show_log=False)
        _paddle_last_used = _now()
        return _paddle_engine, None
    except Exception as exc:  # pragma: no cover - runtime safeguard
        return None, f"PaddleOCR init failed: {exc}"


def _tesseract_ready() -> Tuple[bool, Optional[str]]:
    try:
        import pytesseract  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        return False, f"Tesseract unavailable: {exc}"
    try:
        _ = pytesseract.get_tesseract_version()
        return True, None
    except Exception as exc:
        return False, f"Tesseract unavailable: {exc}"


def convert_pdf_to_images(pdf_path: Path) -> List[Image.Image]:
    images = []
    try:
        from pdf2image import convert_from_path
        # Poppler must be installed in system for this to work
        # In docker: apt-get install -y poppler-utils
        pages = convert_from_path(str(pdf_path), dpi=200)
        for page in pages:
            images.append(page)
    except Exception as e:
        return [] # Changed from f"[Tesseract Error: {e}]" to [] to match return type
    return images


def _ocr_tesseract(img: Image.Image) -> str:
    try:
        import pytesseract
        text = pytesseract.image_to_string(img)
        return text.strip()
    except Exception as e:
        return f"[Tesseract Error: {e}]"


def run_paddle_ocr(image_path: Path) -> Tuple[Optional[str], Optional[str]]:
    engine, err = get_paddle()
    if err:
        return None, err
    try:
        result = engine.ocr(str(image_path), cls=True)
        lines = []
        for page in result:
            for line in page:
                if line and len(line) > 1 and line[1]:
                    lines.append(line[1][0])
        text = "\n".join(lines).strip()
        return text if text else None, None
    except Exception as exc:  # pragma: no cover - runtime safeguard
        return None, f"PaddleOCR failed: {exc}"


def run_tesseract_ocr(image_path: Path) -> Tuple[Optional[str], Optional[str]]:
    ok, err = _tesseract_ready()
    if not ok:
        return None, err
    try:
        from PIL import Image  # type: ignore
        import pytesseract  # type: ignore
        with Image.open(image_path) as img:
            text = pytesseract.image_to_string(img)
        text = (text or "").strip()
        return text if text else None, None
    except Exception as exc:  # pragma: no cover - runtime safeguard
        return None, f"Tesseract OCR failed: {exc}"


def _short_error(message: Optional[str], limit: int = 300) -> Optional[str]:
    """Condense an engine error (Paddle returns whole tracebacks) to its decisive line."""
    if not message:
        return None
    lines = [line.strip() for line in str(message).splitlines() if line.strip()]
    decisive = [line for line in lines if re.search(r"\w*(Error|Exception)\s*:", line)]
    operator = [line for line in lines if line.startswith("[operator")]
    picked = " ".join((decisive[-1:] or lines[:1]) + operator[-1:])
    prefix = lines[0].split(":", 1)[0] if ":" in lines[0] else ""
    if prefix and not picked.startswith(prefix):
        picked = f"{prefix}: {picked}"
    return picked[:limit]


def _run_image_ocr_with_meta(path_obj: Path, engine: str) -> Tuple[Optional[str], Optional[str], Dict[str, Any]]:
    """Run the configured image engine with Tesseract as fallback; report which engine produced text."""
    if engine != "paddle":
        text, err = run_tesseract_ocr(path_obj)
        return text, err, {"method": "tesseract", "engine": "tesseract" if text else None}

    text, paddle_err = run_paddle_ocr(path_obj)
    if text:
        return text, None, {"method": "paddle", "engine": "paddle"}

    paddle_short = _short_error(paddle_err) or "PaddleOCR returned no text"
    logger.warning("PaddleOCR failed for %s; trying Tesseract fallback: %s", path_obj.name, paddle_short)
    fallback_text, fallback_err = run_tesseract_ocr(path_obj)
    if fallback_text:
        return fallback_text, None, {"method": "tesseract_fallback", "engine": "tesseract", "paddle_error": paddle_short}

    errors = [message for message in (paddle_short, _short_error(fallback_err)) if message]
    return None, "; ".join(errors) or "OCR produced no text", {"method": "paddle", "engine": None, "paddle_error": paddle_short}


def _run_image_ocr(path_obj: Path, engine: str) -> Tuple[Optional[str], Optional[str], str]:
    """Run the configured image engine, preserving Tesseract as a safe fallback."""
    text, err, meta = _run_image_ocr_with_meta(path_obj, engine)
    return text, err, meta["method"]


def _extract_pdf_text(path_obj: Path) -> Tuple[Optional[str], Optional[str]]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        return None, f"PDF reader unavailable: {exc}"
    try:
        reader = PdfReader(str(path_obj))
        chunks = []
        for page in reader.pages:
            chunks.append(page.extract_text() or "")
        text = "\n".join(chunks).strip()
        return text if text else None, None
    except Exception as exc:
        return None, f"PDF text extraction failed: {exc}"


def _extract_docx_text(path_obj: Path) -> Tuple[Optional[str], Optional[str]]:
    try:
        from docx import Document  # type: ignore
    except Exception as exc:  # pragma: no cover - optional dependency
        return None, f"DOCX reader unavailable: {exc}"
    try:
        doc = Document(str(path_obj))
        chunks: List[str] = []
        for para in doc.paragraphs:
            txt = (para.text or "").strip()
            if txt:
                chunks.append(txt)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    txt = (cell.text or "").strip()
                    if txt:
                        chunks.append(txt)
        text = "\n".join(chunks).strip()
        return text if text else None, None
    except Exception as exc:
        return None, f"DOCX extraction failed: {exc}"


def _maybe_ocr_pdf(path_obj: Path) -> Tuple[Optional[str], Optional[str]]:
    text, err, _info = _maybe_ocr_pdf_with_meta(path_obj)
    return text, err


def _maybe_ocr_pdf_with_meta(path_obj: Path) -> Tuple[Optional[str], Optional[str], Dict[str, Any]]:
    """OCR rendered PDF pages; ``info`` records the engine that produced text and any Paddle error."""
    info: Dict[str, Any] = {"engine": None}
    # Try pymupdf (fitz) first, then fall back to pdf2image
    try:
        import fitz  # pymupdf
        import tempfile
        import os as _os
        
        doc = fitz.open(str(path_obj))
        all_text = []
        for page_num in range(min(3, len(doc))):  # Limit to first 3 pages
            page = doc[page_num]
            pix = page.get_pixmap(dpi=150)
            
            # Create temp file path without keeping it open (Windows compat)
            tmp_fd, tmp_path = tempfile.mkstemp(suffix=".png")
            _os.close(tmp_fd)  # Close immediately so pix.save can write
            
            try:
                pix.save(tmp_path)
                
                engine = _resolve_engine()
                text, err, page_meta = _run_image_ocr_with_meta(Path(tmp_path), engine)
                if page_meta.get("paddle_error"):
                    info["paddle_error"] = page_meta["paddle_error"]

                if text:
                    all_text.append(text)
                    info["engine"] = info["engine"] or page_meta.get("engine")
            finally:
                try:
                    _os.unlink(tmp_path)
                except Exception:
                    pass
        
        doc.close()
        combined = "\n".join(all_text).strip()
        return (combined if combined else None), None, info
    except ImportError:
        pass  # Fall through to pdf2image
    except Exception as exc:
        return None, f"PDF OCR via pymupdf failed: {exc}", info

    # Fallback to pdf2image (requires poppler)
    try:
        from pdf2image import convert_from_path  # type: ignore
    except Exception:
        return None, "PDF OCR unavailable (install pdf2image/poppler or pymupdf).", info
    try:
        images = convert_from_path(str(path_obj), first_page=1, last_page=1)
        if not images:
            return None, "PDF OCR failed: no pages rendered.", info
        image = images[0]
        try:
            import pytesseract  # type: ignore
        except Exception as exc:
            return None, f"Tesseract unavailable: {exc}", info
        text = pytesseract.image_to_string(image)
        text = (text or "").strip()
        if text:
            info["engine"] = "tesseract"
        return text if text else None, None, info
    except Exception as exc:
        return None, f"PDF OCR failed: {exc}", info


def extract_text_with_meta(image_path: str, min_chars: int | None = None) -> Tuple[Optional[str], Optional[str], Dict[str, Any]]:
    """
    Try text extraction first; OCR only if text is too short.
    Returns (text, error, meta).
    """
    path_obj = Path(image_path)
    if not path_obj.exists():
        return None, f"File not found: {image_path}", {"ocr_used": False, "method": "missing"}

    min_chars = _OCR_MIN_TEXT_CHARS if min_chars is None else min_chars
    suffix = path_obj.suffix.lower()

    if suffix in {".txt", ".md", ".log", ".json"}:
        try:
            text = path_obj.read_text(encoding="utf-8", errors="ignore").strip()
            return (text if text else None), None, {"ocr_used": False, "method": "text", "engine": "text"}
        except Exception as exc:
            return None, f"Text read failed: {exc}", {"ocr_used": False, "method": "text", "engine": None}

    if suffix == ".docx":
        text, err = _extract_docx_text(path_obj)
        if text:
            return text, None, {"ocr_used": False, "method": "docx", "engine": "docx"}
        return None, err or "DOCX extraction produced no content.", {"ocr_used": False, "method": "docx", "engine": None}

    if suffix == ".doc":
        return (
            None,
            "Legacy .doc files are not supported. Please upload PDF, image, or .docx.",
            {"ocr_used": False, "method": "doc"},
        )

    if suffix == ".pdf":
        text, err = _extract_pdf_text(path_obj)
        if text and len(text) >= min_chars:
            return text, None, {"ocr_used": False, "method": "pdf", "engine": "pdf"}
        ocr_text, ocr_err, ocr_info = _maybe_ocr_pdf_with_meta(path_obj)
        if ocr_text:
            return ocr_text, None, {"ocr_used": True, "method": "pdf_ocr", **ocr_info}
        meta: Dict[str, Any] = {"ocr_used": False, "method": "pdf", "engine": "pdf" if text else None}
        if ocr_info.get("paddle_error"):
            meta["paddle_error"] = ocr_info["paddle_error"]
        return text, ocr_err or err or "PDF text extraction produced no content.", meta

    engine = _resolve_engine()
    text, err, image_meta = _run_image_ocr_with_meta(path_obj, engine)

    if text:
        log_action("ocr_call", f"engine={image_meta['engine']} method={image_meta['method']} path={image_path}")
        return text, None, {"ocr_used": True, **image_meta}
    return None, err or "OCR produced no text", {"ocr_used": True, **image_meta}


def extract_text(image_path: str) -> Tuple[Optional[str], Optional[str]]:
    text, err, _meta = extract_text_with_meta(image_path)
    return text, err


def _probe_engine(name: str) -> Dict[str, Any]:
    """OCR the bundled health image with one engine and check the expected token comes back."""
    if name == "paddle" and find_spec("paddleocr") is None:
        return {"ok": False, "error": "PaddleOCR package missing"}
    runner = run_paddle_ocr if name == "paddle" else run_tesseract_ocr
    try:
        text, err = runner(_HEALTH_IMAGE)
    except Exception as exc:  # pragma: no cover - runtime safeguard
        text, err = None, f"{name} OCR raised: {exc}"
    if text and _HEALTH_EXPECTED_TOKEN in text:
        return {"ok": True, "error": None}
    return {"ok": False, "error": _short_error(err) or f"{name} OCR did not read the test image (got {text!r})"}


def health_report(force: bool = False) -> Dict[str, Any]:
    """Run real OCR on a bundled test image; cached for OCR_HEALTH_TTL_SEC seconds."""
    global _health_cache
    if not force and _health_cache and (_now() - _health_cache[0]) < _HEALTH_TTL_SEC:
        return _health_cache[1]
    configured = _resolve_engine()
    engines: Dict[str, Dict[str, Any]] = {}
    if configured == "paddle":
        engines["paddle"] = _probe_engine("paddle")
    engines["tesseract"] = _probe_engine("tesseract")
    configured_ok = engines.get(configured, {}).get("ok", False)
    if configured_ok:
        active = configured
        detail = f"{'PaddleOCR' if configured == 'paddle' else 'Tesseract'} read the test image"
    elif configured == "paddle" and engines["tesseract"]["ok"]:
        active = "tesseract"
        detail = f"{engines['paddle']['error']}; Tesseract fallback read the test image"
    else:
        active = None
        detail = "; ".join(f"{name} failed: {probe['error']}" for name, probe in engines.items())
    report = {
        "ok": configured_ok,
        "detail": detail,
        "configured_engine": configured,
        "active_engine": active,
        "engines": engines,
    }
    _health_cache = (_now(), report)
    return report


def health() -> Tuple[bool, str]:
    """True only when the configured engine actually reads the bundled test image."""
    report = health_report()
    return report["ok"], report["detail"]
